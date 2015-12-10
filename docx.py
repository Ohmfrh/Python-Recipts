from docx import Document
from docx.shared import Mm

__author__ = 'daniel'

document = Document()
section = document.sections

section.page_width = Mm(80)
section.page_height = Mm(297)


paragraph = document.add_paragraph('Lorem ipsum dolor sit amet.')
document.add_heading('The REAL meaning of the universe')
document.add_heading('The role of dolphins', level=2)


document.save('test.docx')
